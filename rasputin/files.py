"""File parsing for document review."""

from pathlib import Path

# Hard cap on raw file size, checked before reading, to avoid loading a huge
# file into memory (OOM). Separate from the per-file character cap, which
# bounds how much extracted text is actually sent to the models.
MAX_FILE_BYTES = 50 * 1024 * 1024  # 50 MB


def parse_files(paths: list[str], max_chars: int | None = None) -> str:
    """Parse multiple files and return concatenated text.

    ``max_chars`` truncates each file's extracted text (with a marker) to bound
    token usage and cost; ``None`` means no truncation.
    """
    parts = []
    for path in paths:
        try:
            text = parse_file(path)
            if max_chars is not None and len(text) > max_chars:
                omitted = len(text) - max_chars
                text = text[:max_chars] + f"\n\n[TRUNCATED: {omitted} characters omitted]"
            parts.append(f"=== FILE: {Path(path).name} ===\n{text}")
        except Exception as e:
            parts.append(f"=== FILE: {Path(path).name} [ERROR: {e}] ===")
    return "\n\n".join(parts)


def parse_file(path: str) -> str:
    """Parse a single file based on extension."""
    p = Path(path)
    if p.exists() and p.stat().st_size > MAX_FILE_BYTES:
        raise ValueError(
            f"File too large: {p.stat().st_size} bytes (max {MAX_FILE_BYTES})"
        )
    ext = p.suffix.lower()
    parsers = {
        ".pdf": _parse_pdf,
        ".docx": _parse_docx,
        ".pptx": _parse_pptx,
        ".xlsx": _parse_xlsx,
        ".csv": _parse_csv,
        ".txt": _parse_text,
        ".md": _parse_text,
        ".json": _parse_text,
        ".py": _parse_text,
        ".js": _parse_text,
        ".ts": _parse_text,
    }
    parser = parsers.get(ext)
    if not parser:
        raise ValueError(f"Unsupported format: {ext}")
    return parser(path)


def _parse_pdf(path: str) -> str:
    import pdfplumber

    parts = []
    with pdfplumber.open(path) as pdf:
        for i, page in enumerate(pdf.pages, 1):
            text = page.extract_text() or ""
            tables = page.extract_tables()
            table_text = ""
            if tables:
                for table in tables:
                    rows = [
                        " | ".join(str(cell or "") for cell in row) for row in table
                    ]
                    table_text += "\n".join(rows) + "\n"
            parts.append(f"--- Page {i} ---\n{text}\n{table_text}".strip())
    return "\n\n".join(parts)


def _parse_docx(path: str) -> str:
    from docx import Document

    doc = Document(path)
    parts = []
    for para in doc.paragraphs:
        if para.text.strip():
            prefix = (
                f"[{para.style.name}] "
                if para.style.name.startswith("Heading")
                else ""
            )
            parts.append(f"{prefix}{para.text}")
    for table in doc.tables:
        rows = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            rows.append(" | ".join(cells))
        parts.append("\n".join(rows))
    return "\n\n".join(parts)


def _parse_pptx(path: str) -> str:
    from pptx import Presentation

    prs = Presentation(path)
    parts = []
    for i, slide in enumerate(prs.slides, 1):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = para.text.strip()
                    if text:
                        texts.append(text)
        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes_text = slide.notes_slide.notes_text_frame.text.strip()
            if notes_text:
                notes = f"\n  Notes: {notes_text}"
        slide_text = "\n  ".join(texts)
        parts.append(f"--- Slide {i} ---\n  {slide_text}{notes}")
    return "\n\n".join(parts)


def _parse_xlsx(path: str) -> str:
    from openpyxl import load_workbook

    # read_only streams rows instead of loading the whole workbook, limiting
    # blow-up from oversized (or maliciously crafted) spreadsheets.
    wb = load_workbook(path, data_only=True, read_only=True)
    parts = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = []
        for row in ws.iter_rows(values_only=True):
            cells = [str(cell) if cell is not None else "" for cell in row]
            if any(cells):
                rows.append(" | ".join(cells))
        if rows:
            parts.append(f"--- Sheet: {sheet_name} ---\n" + "\n".join(rows))
    return "\n\n".join(parts)


def _parse_csv(path: str) -> str:
    import csv

    with open(path, newline="", encoding="utf-8") as f:
        reader = csv.reader(f)
        rows = [" | ".join(row) for row in reader]
    return "\n".join(rows)


def _parse_text(path: str) -> str:
    return Path(path).read_text(encoding="utf-8")
