from pathlib import Path

import pytest

FIXTURES = Path(__file__).parent / "fixtures"


def test_parse_text_file():
    from rasputin.files import parse_file

    text = parse_file(str(FIXTURES / "sample.txt"))
    assert "test document" in text


def test_parse_unsupported_format():
    from rasputin.files import parse_file

    with pytest.raises(ValueError, match="Unsupported format"):
        parse_file("/tmp/fake.xyz")


def test_parse_files_multiple():
    from rasputin.files import parse_files

    result = parse_files([str(FIXTURES / "sample.txt")])
    assert "FILE: sample.txt" in result
    assert "test document" in result


def test_parse_files_handles_errors():
    from rasputin.files import parse_files

    result = parse_files(["/nonexistent/file.txt"])
    assert "ERROR" in result


def test_parse_csv(tmp_path):
    from rasputin.files import parse_file

    p = tmp_path / "data.csv"
    p.write_text("name,score\nAlice,10\n", encoding="utf-8")
    out = parse_file(str(p))
    assert "name | score" in out
    assert "Alice | 10" in out


def test_parse_docx(tmp_path):
    from docx import Document

    from rasputin.files import parse_file

    p = tmp_path / "doc.docx"
    doc = Document()
    doc.add_paragraph("Hello from docx")
    doc.save(str(p))
    out = parse_file(str(p))
    assert "Hello from docx" in out


def test_parse_xlsx(tmp_path):
    from openpyxl import Workbook

    from rasputin.files import parse_file

    p = tmp_path / "book.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["col1", "col2"])
    ws.append([1, 2])
    wb.save(str(p))
    out = parse_file(str(p))
    assert "col1 | col2" in out
    assert "1 | 2" in out


def test_parse_pptx(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches

    from rasputin.files import parse_file

    p = tmp_path / "deck.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    box.text_frame.text = "Slide content here"
    prs.save(str(p))
    out = parse_file(str(p))
    assert "Slide content here" in out


def test_parse_files_truncates_to_max_chars(tmp_path):
    from rasputin.files import parse_files

    p = tmp_path / "big.txt"
    p.write_text("x" * 1000, encoding="utf-8")
    out = parse_files([str(p)], max_chars=100)
    assert "[TRUNCATED:" in out
    assert "900 characters omitted" in out


def test_parse_file_rejects_oversize(tmp_path, monkeypatch):
    import rasputin.files as files_mod

    monkeypatch.setattr(files_mod, "MAX_FILE_BYTES", 10)
    p = tmp_path / "toobig.txt"
    p.write_text("x" * 100, encoding="utf-8")
    with pytest.raises(ValueError, match="too large"):
        files_mod.parse_file(str(p))
